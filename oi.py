import sys
from numpy import transpose
import requests
import json
import pandas as pd
import time

sym = 'NIFTY'
exp_date = '02-Jun-2022'

requestTimeout = 10

#for xlwt
'''
import xlwt
from xlwt import Workbook
wb = Workbook()
sheet1 = wb.add_sheet('Sheet3', cell_overwrite_ok=True)
rownum = 0
colnum = 0
'''

#for openpyxl
'''
import openpyxl
my_wb = openpyxl.Workbook()
my_sheet = my_wb.active
rownum = 1
colnum = 1
'''

import xlwings as xw
from xlwings.constants import AutoFillType
import xlsxwriter as wr 
#from xlwings import range as xlrange

file = xw.Book('oi.xlsx')
sh1 = file.sheets("Sheet1")
sh2 = file.sheets("Sheet2")
rownum = 1
colnum = 0
atmColumn = None
prevAtmColumn = None

a=1

cookie = '_ga=GA1.2.176340240.1650143938; arp_scroll_position=141; nsit=UQyyffPlSXC1GL91-GLC3IYH; AKA_A2=A; nseappid=eyJhbGciOiJIUzI1NiIsInR5cCI6IkpXVCJ9.eyJpc3MiOiJhcGkubnNlIiwiYXVkIjoiYXBpLm5zZSIsImlhdCI6MTY1NDEyNDIyNSwiZXhwIjoxNjU0MTI3ODI1fQ.F1DrR8JwHLT7P8147T9iFjqhh6sKLr2or7bxVPQgj-w; ak_bmsc=77B1AD9E26302C5A7B32E87DB397437E~000000000000000000000000000000~YAAQBNcLF2kP8/OAAQAAnqV8IQ+EXPz1ZgeogEJ5QRtC57pKgdPZvUbZP3G2aQU3wtyJbszOXRz3e/CKA0fZ9NQeWtyBPwDuCI68AVoKS7+Vs4W3nNGNFqMZYEoHGJHfhcn1CJdUZANHJ0LVrn2LccdJ6yo8zB6tbQLTNliKCBGngV6eeMAKvzVBmHCSzzGIBI9ahyNV92emIx7Se2QhVCEGnX0dopbuO0czAXwd6M9O99oPlzGkzQPM+KovHJ5iEFVFe8592tXnB1ewdNyxSb+YcRTzHcwk9g8X00UwmbBxhscLLU6Qqt0X+ZfcAnTNckroDyXAX/ROVAxTx44n/WVZvOnH0UakAb8PgbtAWEoODXfc9uiLajfiVidS1II0XGOQUYy/X8DziDUm1yZeHH0D4gq3kZ9oUb6QU1RVMIMNhFcgJ2mpzXgyTa7Ar6GVeK8toDui93AlsP+ojPtUOobypNZps24aN4B/wMWywdHN; RT="z=1&dm=nseindia.com&si=ca6b0f71-5a98-4dbb-aa81-c35b6aeea104&ss=l3w599rp&sl=0&tt=0&bcn=%2F%2F684d0d4c.akstat.io%2F"; bm_sv=BE70E68C35CEA860AA602FC6FCA3644E~YAAQBNcLF20P8/OAAQAADKh8IQ8WfSJCZArZx5Aawh8yz8PIz/8cYkaRZ4mlXbVxvZQ6bLipobOsSKsgNYwrqHJSHVSK8Z/3ZRj+LOT9pOLkGVL38G9X1+Gpp8QnvBoGuQ2ZkfpZaTpYh3RokYavvFq6HSLfSuWnOwzIn9cPTrhfPQ+XYwkUqnUBSdF/zkAUrCAtIstb5xrs9xUWN2NPRGAIzakiDuJuREWpUvodSSSlu628qlE0WyuOdfJGDNFZcwo=~1' 


# This function is for countdown timer    
def timeCountdown(t):
    while t:
        mins, secs = divmod(t, 60)
        timeformat = '{:02d}:{:02d}'.format(mins, secs)
        print("Time Remaining : " + timeformat, end='\r')
        time.sleep(1)
        t -= 1
    print('\n')


# This function pulls the option chain data for selected symbol
def pullOptionChain(symbol):
    try:
        if symbol == 'NIFTY' or symbol == 'BANKNIFTY':
            url = 'https://www.nseindia.com/api/option-chain-indices?symbol=' + symbol
        else:
            url = 'https://www.nseindia.com/api/option-chain-equities?symbol=' + symbol 
        headers = {
          "accept" : "*/*",
          "accept-encoding" : "gzip, deflate, br",
          "accept-language" : "en-US,en;q=0.9",
          "cookie" : cookie,
          "referer" : "https://www.nseindia.com/option-chain", 
          "user-agent" : "Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/101.0.4951.41 Safari/537.36"
        }
        response = requests.get(url, headers = headers, timeout = requestTimeout).text
        data = json.loads(response)
        return data
    except Exception as e:
        print(e)
        print("\nUpdate the cookies from NSE website".upper())
        sys.exit(1)


# This function writes the OI data to Excel sheet
def writeOIDataToExcel(oi, colnum, underlying_price, atmColumn):
    try:
        global rownum
        global prevAtmColumn
        if prevAtmColumn != None:
            colname = wr.utility.xl_col_to_name(prevAtmColumn)
            cell = colname + str(1)
            sh1.range(cell).clear()
        if atmColumn != None:
            #colname = wr.utility.xl_col_to_name(colnum)
            #cell = colname + str(rownum)
            colname = wr.utility.xl_col_to_name(atmColumn)
            cell = colname + str(1)
            sh1.autofit('columns')
            #sh1.range(cell).api.autofill(sh1.range(cell), AutoFillType.xlFillDefault)
            sh1.range(cell).color = (10, 220, 10)
            sh1.range(cell).font.bold = True
            sh1.range(cell).font.size = 16
            sh1.range(cell).value=underlying_price
            sh1.range(cell).row_height = 30
            prevAtmColumn = atmColumn
            #rownum = rownum + 1
            
        if rownum == 2 and colnum == 0:
            for x in oi.values():
                #print(str(x['strikeprice']))
                colname = wr.utility.xl_col_to_name(colnum)
                #print(colname)
                #print(colname + str(rownum))
                cell = colname + str(rownum)
                sh1.autofit('columns')
                sh1.range(cell).color = (255, 100, 55)
                sh1.range(cell).font.bold = True
                sh1.range(cell).font.size = 13
                #sh1.range(cell).api.Borders(9).LineStyle = 1
                #sh1.range(cell).api.Borders(9).Weight = 3
                #sh1.range(cell).api.HorizontalAlignment = -4108
                sh1.range(cell).value=str(x['strikeprice'])
                colnum = colnum + 1 
        elif rownum > 2: 
            print("Writing oi data to excel")
            colnum = 0
            for x in oi.values():
                #print(str(x['openInterest']))
                colname = wr.utility.xl_col_to_name(colnum)
                cell = colname + str(rownum)
                sh1.autofit('columns')
                sh1.range(cell).font.size = 13
                #print(colname)
                #print(colname + str(rownum)) 
                #sh1.range(cell).api.Borders(9).LineStyle = 1
                #sh1.range(cell).api.Borders(9).Weight = 3
                #sh1.range(cell).api.HorizontalAlignment = -4108
                sh1.range(cell).value=str(x['openInterest'])
                colnum = colnum + 1
    except Exception as e:
        print("Writing to Excel Failed with Exception : " + str(e))
        sys.exit(1)


# This function gets the closest price to the underlying price which is the ATM price
def closestValueInList(strike_prices, underlying_price):
    try:
        return (min(strike_prices, key=lambda x:abs(x-underlying_price)))
    except Exception as e:
        print("ATM Price - Getting closest value in strikePrices list failed with Exception : " + str(e))
        sys.exit(1)


# This function gets the column number which has the ATM price
def getColumnIndexInExcelHavingValue(sh1, atmPrice, strikePriceList):
    try:
        for col in range(0, len(strikePriceList)):
            cell = sh1.range('{}{}'.format(wr.utility.xl_col_to_name(col), 2)).value
            if cell == atmPrice:
                return col
    except Exception as e:
        print("Getting column index in excel having value failed with Exception : " + str(e))
        sys.exit(1)


# hide few columns in Excel that are far from the ATM Price
def hideColumnInExcel(sh1, atmPricecolnum):
    try:
        toHideColumns = 10
        for col in range(0, atmPricecolnum - toHideColumns):
            sh1.range('{}{}'.format(wr.utility.xl_col_to_name(col), 2)).api.EntireColumn.Hidden = True
            
    except Exception as e:
        print("Hiding columns in excel failed with Exception : " + str(e))
        sys.exit(1)


# Loop starts here
while True:
    try: 
        data = pullOptionChain(sym)
        if len(data) == 0:
            print("No data pulled from NSE site, check the Expiry date")
            break
        else:
            print("Got Response from NSE")
            #print(json.dumps(data, indent= 4))
            exp_list = data['records']['expiryDates']
            underlying_value = data['records']['underlyingValue']
            strikePriceList = data['records']['strikePrices']
            atmPrice = closestValueInList(data['records']['strikePrices'], underlying_value)
            print("symbol : " + sym.upper() + "\nUnderlying Price : " + str(underlying_value) + "\nATM Price : " + str(atmPrice))
            #print(data['records']['strikePrices'])
            ce = {}
            pe = {}
            n = 0
            m = 0
            oi = {}
            x = 0
            """
            for i in data['records']['data'] :
                if i['expiryDate'] == exp_date:
                    try: 
                        ce[n] = i['CE']
                        n = n+1
                    except:
                        pass
                    try:
                        pe[m] = i['PE']
                        m = m+1
                    except:
                        pass
            """

            """
            n = 0 
            for i in data['records']['data']:
                if i['expiryDate'] == exp_date:
                    try:
                        oi[n] = { 
                            "changeinOpenInterest_ce" : i["CE"]["changeinOpenInterest"],
                            "pchangeinOpenInterest_ce" : i["CE"]["pchangeinOpenInterest"],
                            "impliedVolatility" : i["CE"]["impliedVolatility"],
                            "change_ce" : i["CE"]["change"],
                            "lastPrice_ce" : i["CE"]["lastPrice"],
                            "bidprice_ce" : i["CE"]["bidprice"],
                            "askPrice_ce" : i["CE"]["askPrice"],
                            "oi_ce" : i['CE']["openInterest"] / 2,

                            "strikeprice" : i['strikePrice'], 

                            "oi_pe" : i["PE"]["openInterest"] / 2,
                            "bidprice_pe" : i["PE"]["bidprice"],
                            "askPrice_pe" : i["PE"]["askPrice"],
                            "lastPrice_pe" : i["PE"]["lastPrice"],
                            "change_pe" : i["PE"]["change"],
                            "impliedVolatility" : i["PE"]["impliedVolatility"],
                            "changeinOpenInterest_pe" : i["PE"]["changeinOpenInterest"],
                            "pchangeinOpenInterest_pe" : i["PE"]["pchangeinOpenInterest"],
                        }
                        n = n+1
                    except:
                        pass
            """
            n = 0
            for i in data['records']['data']:
                if i['expiryDate'] == exp_date:
                    try:
                        oi[n] = { 
                            "strikeprice" : i['strikePrice'], 
                            "openInterest" : "|".join([str(i['CE']["openInterest"] / 2), str(i['PE']["openInterest"] / 2)])
                        }
                        n = n+1
                    except:
                        pass  
            
            if len(oi) == 0:
                print("OpenInterest data not available, check the Expiry date. \nAvailable Expiry Dates : ")
                print(*exp_list, sep='\n')
                sys.exit(1)
            #print("OPENINTEREST") 
            #print(json.dumps(oi, indent = 4))
            #print(type(oi[0]['openInterest']))

            #xlwings ( option - 1 )  
            writeOIDataToExcel(oi, colnum, underlying_value, atmColumn)
            rownum = rownum + 1  
            atmColumn = getColumnIndexInExcelHavingValue(sh1, atmPrice, strikePriceList)
            #print("Column having ATM Value : " + str(atmColumn))
            #TO-DO : Hide columns far from the ATM price  
            #hideColumnInExcel(sh1, atmColumn)
            
            #xlwings ( options - 2)
            '''
            #using dataframe  - working
            df = pd.DataFrame.from_dict(oi).transpose()
            print(df)
            if a == 1 :
                oi_df = df[['strikeprice', 'oi']]
            else:
                oi_df = df[['oi']]

            print(oi_df.transpose())
            data = oi_df.transpose()
            sh1.range("A"+str(a)).value = data

            if a == 1:
                a = a + 3
            else: a = a + 2
            '''
    
        
            '''
            #xlwt - only works with .xls(old format excel)
            # This is working but excel values are not getting refreshed . Have to close and open again to see new writes
            if rownum == 0 and colnum == 0:
                for x in oi.values():
                    sheet1.write(rownum, colnum, x['strikeprice'])
                    colnum = colnum + 1
            else: 
                colnum = 0
                for x in oi.values():
                    sheet1.write(rownum, colnum, str(x['oi']))
                    colnum = colnum + 1

            wb.save('book1.xls')
            rownum = rownum + 1
            '''
            '''
            #openpyxl
            # This is working but excel values are not getting refreshed . Have to close and open again to see new writes
            if rownum == 1 and colnum == 1:
                for x in oi.values():
                    my_sheet.cell(row=rownum, column=colnum).value=x['strikeprice']
                    colnum = colnum + 1
            else: 
                colnum = 1
                for x in oi.values():
                    my_sheet.cell(row=rownum, column=colnum).value=str(x['oi'])
                    colnum = colnum + 1
            my_wb.save("book3.xlsx") 
            rownum = rownum + 1       
            '''       
            
            timeCountdown(6)
            #time.sleep(180)
    
    
            '''
            oi_df = pd.DataFrame.from_dict(oi).transpose()

            print(oi_df)

            if a == 1:
                df_oi = oi_df[['strikeprice',  'oi']]
            else:
                df_oi = oi_df[['oi']]

            print(df_oi.transpose())
            data = df_oi.transpose() 
            data.to_excel('book1.xlsx', sheet_name='Sheet1", index=False')   
            '''
    except Exception as e:
        print(e)
        print("End of Script")
        #print("\nUpdate the cookies from NSE website".upper())
        sys.exit(1)
